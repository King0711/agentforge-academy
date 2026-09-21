"""
Step 4 of the Contract Clause Extractor (part 2) - exports the clause
review to a formatted Word document. No AI or API key needed - this is
pure document assembly, fully testable by actually building and
re-opening a real .docx file.
"""

from docx import Document
from docx.shared import RGBColor

RISK_COLORS = {
    "red": RGBColor(0xC0, 0x00, 0x00),
    "yellow": RGBColor(0xB8, 0x86, 0x0B),
    "green": RGBColor(0x00, 0x80, 0x00),
    "unclear": RGBColor(0x80, 0x80, 0x80),
}


def export_word(clauses, output="contract_review.docx"):
    """
    Builds a Word document: a title, then one section per clause with a
    color-coded risk line, a plain-English summary, and the exact quote.

        export_word(clauses, output="review.docx")
    """
    document = Document()
    document.add_heading("Contract Clause Review", level=0)

    for clause_type, data in clauses.items():
        document.add_heading(clause_type, level=1)

        risk = data.get("risk", "unclear")
        risk_paragraph = document.add_paragraph()
        risk_run = risk_paragraph.add_run(f"Risk: {risk.upper()}")
        risk_run.font.color.rgb = RISK_COLORS.get(risk, RISK_COLORS["unclear"])
        risk_run.bold = True

        if data.get("verified") is False:
            warning = document.add_paragraph()
            warning.add_run("Quote could not be verified against the source contract - check by hand.").italic = True

        document.add_paragraph(data.get("summary", ""))
        document.add_paragraph(data.get("exact_text", ""))

    document.save(output)
    return output


if __name__ == "__main__":
    # Run this file on its own to check the Word export works - a real
    # .docx file built, re-opened, and verified, then cleaned up:
    #     python export.py
    import os

    clauses = {
        "Payment Terms": {"summary": "Net 30 payment terms.", "exact_text": "Payment due within 30 days.", "risk": "green", "verified": True},
        "Non-Compete": {"summary": "5-year non-compete.", "exact_text": "Shall not compete for 5 years.", "risk": "red", "verified": False},
    }

    path = export_word(clauses, output="test_review.docx")
    assert os.path.exists(path)

    reopened = Document(path)
    all_text = "\n".join(p.text for p in reopened.paragraphs)
    headings = [p.text for p in reopened.paragraphs if p.style.name.startswith("Heading")]

    assert "Contract Clause Review" in all_text
    assert "Payment Terms" in headings
    assert "Non-Compete" in headings
    assert "RISK: GREEN" in all_text.upper() or "RISK: RED" in all_text.upper()
    assert "could not be verified" in all_text
    print("Re-opened the saved document and confirmed the title, both clause headings,")
    print("a risk line, and the unverified-quote warning are all present.")

    os.remove(path)
    print("\nAll checks passed.")
